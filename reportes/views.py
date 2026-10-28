from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import render
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from cuentas.models import Rol
from cuentas.permisos import rol_requerido
from instrumento.forms import SEMESTRES

from .services import filtrar, resumen

SOLO_DOCENTES = (Rol.DOCENTE, Rol.ADMIN)


def _parametros(request):
    semestre = request.GET.get("semestre", "")
    if semestre and (not semestre.isdigit() or not 0 <= int(semestre) <= 10):
        semestre = ""                      # valor inválido: se ignora el filtro
    return semestre, request.GET.get("q", "").strip()[:60]


@rol_requerido(*SOLO_DOCENTES)
def dashboard(request):
    semestre, q = _parametros(request)
    qs = filtrar(semestre, q)
    datos = resumen(qs)
    contexto = {
        **datos, "semestre": semestre, "q": q, "semestres": SEMESTRES[1:],
        "recientes": qs[:50],
        "muestra_pequena": datos["total_estudiantes"] < settings.MUESTRA_MINIMA,
        "querystring": request.GET.urlencode(),
    }
    # HTMX pide solo el panel; una carga normal devuelve la página completa.
    plantilla = "reportes/_panel.html" if request.headers.get("HX-Request") else "reportes/dashboard.html"
    return render(request, plantilla, contexto)


@rol_requerido(*SOLO_DOCENTES)
def exportar_excel(request):
    semestre, q = _parametros(request)
    qs = filtrar(semestre, q)
    datos = resumen(qs)

    wb = Workbook()
    negrita = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="2F5597")

    def encabezado(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = negrita, relleno

    ws = wb.active
    ws.title = "Resumen"
    encabezado(ws, ["Indicador", "Valor"])
    ws.append(["Filtro de semestre", semestre or "Todos"])
    ws.append(["Estudiantes", datos["total_estudiantes"]])
    ws.append(["Intentos", datos["total_intentos"]])
    for nivel, n in datos["distribucion"].items():
        ws.append([f"Riesgo {nivel}", n])

    ws = wb.create_sheet("Variables")
    encabezado(ws, ["Código", "Variable", "Promedio", "Semáforo"])
    for v in datos["variables"]:
        ws.append([v["codigo"], v["nombre"], round(v["promedio"], 2), v["color"]])

    ws = wb.create_sheet("Intentos")
    encabezado(ws, ["Estudiante", "Semestre", "Puntaje", "Nivel", "Versión de rangos", "Fecha"])
    for i in qs:
        ws.append([i.usuario.get_full_name() or i.usuario.username, i.semestre, i.puntaje_total,
                   i.get_nivel_display(), i.config.version, i.fecha.strftime("%Y-%m-%d %H:%M")])
    for hoja in wb:
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = 22

    respuesta = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respuesta["Content-Disposition"] = 'attachment; filename="reporte_alfabetizacion.xlsx"'
    wb.save(respuesta)
    return respuesta
